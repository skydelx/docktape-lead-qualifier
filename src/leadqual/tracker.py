"""Step 5: write the result to the Excel tracker a sales rep opens and skims."""

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from leadqual.models import Result, Route

SHEET_NAME = "Leads"
# (header, column width)
COLUMNS: tuple[tuple[str, int], ...] = (
    ("Route", 16),
    ("Why", 42),
    ("Company", 26),
    ("Website", 24),
    ("Fit", 6),
    ("Data", 9),
    ("Compliance", 34),
    ("Compliance reasoning", 60),
    ("Summary", 60),
    ("Size", 32),
    ("Cloud signal", 38),
    ("HQ", 18),
    ("Contact", 38),
    ("Updated", 17),
)
HEADERS = [header for header, _ in COLUMNS]
COMPANY_COLUMN = HEADERS.index("Company") + 1
WEBSITE_COLUMN = HEADERS.index("Website") + 1
ROUTE_COLOURS = {
    Route.SALES_READY: "C6EFCE",
    Route.REVIEW: "FFEB9C",
    Route.DO_NOT_ENGAGE: "FFC7CE",
    Route.LOW_PRIORITY: "E7E6E6",
}


class TrackerError(RuntimeError):
    """The spreadsheet could not be written."""


def record(result: Result, path: Path) -> int:
    """Add the company's row, or update it if the company is already there. Returns the row."""
    workbook = load_workbook(path) if path.exists() else _new_workbook()
    sheet = workbook[SHEET_NAME]
    row = _existing_row(sheet, result) or sheet.max_row + 1
    for column, value in enumerate(_row_values(result), start=1):
        cell = sheet.cell(row=row, column=column, value=value)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.cell(row=row, column=1).fill = PatternFill(
        "solid", start_color=ROUTE_COLOURS[result.decision.route]
    )
    sheet.auto_filter.ref = sheet.dimensions
    try:
        workbook.save(path)
    except PermissionError as error:
        raise TrackerError(f"{path} is open in another program; close it and run again") from error
    return row


def _new_workbook() -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    for column, (header, width) in enumerate(COLUMNS, start=1):
        sheet.cell(row=1, column=column, value=header).font = Font(bold=True)
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A2"
    return workbook


def _existing_row(sheet: Worksheet, result: Result) -> int | None:
    """One row per company: matched by website domain, or by name when there is no website."""
    domain = _domain(result)
    name = result.company.name.casefold()
    for row in range(2, sheet.max_row + 1):
        website = sheet.cell(row=row, column=WEBSITE_COLUMN).value or ""
        company = str(sheet.cell(row=row, column=COMPANY_COLUMN).value or "")
        if website == domain and (domain or company.lstrip("'").casefold() == name):
            return row
    return None


def _domain(result: Result) -> str:
    return result.company.domain or ""


def _row_values(result: Result) -> list[str | int]:
    contact, research, compliance, fit = (
        result.contact,
        result.research,
        result.compliance,
        result.fit,
    )
    title = f", {contact.job_title}" if contact.job_title else ""
    values: list[str | int] = [
        result.decision.route.value,
        result.decision.reason,
        result.company.name,
        _domain(result),
        fit.total,
        fit.data,
        compliance.flag,
        compliance.reasoning,
        research.summary,
        fit.size_basis,
        fit.cloud_basis,
        compliance.hq_country or research.hq_country or result.company.declared_country or "",
        f"{contact.name}{title} <{contact.email}>",
        result.processed_at.strftime("%Y-%m-%d %H:%M"),
    ]
    return [_safe(value) for value in values]


def _safe(value: str | int) -> str | int:
    """Text from a form or a website must never be run by Excel as a formula."""
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value
