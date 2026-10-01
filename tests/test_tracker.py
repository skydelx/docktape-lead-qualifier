import pytest
from openpyxl import Workbook, load_workbook

from fakes import make_result
from leadqual.models import Route
from leadqual.tracker import HEADERS, ROUTE_COLOURS, TrackerError, record


@pytest.fixture
def path(tmp_path):
    return tmp_path / "leads.xlsx"


def rows(path) -> list[dict]:
    sheet = load_workbook(path)["Leads"]
    header, *body = sheet.iter_rows(values_only=True)
    return [dict(zip(header, row, strict=True)) for row in body]


def test_first_lead_creates_a_spreadsheet_a_sales_rep_can_skim(path):
    row = record(make_result(), path)

    sheet = load_workbook(path)["Leads"]
    assert row == 2
    assert [cell.value for cell in sheet[1]] == HEADERS
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == sheet.dimensions
    assert rows(path) == [
        {
            "Route": "SALES-READY",
            "Why": "Compliance clear; high priority (fit 85)",
            "Company": "Acme Analytics",
            "Website": "acme.io",
            "Fit": 85,
            "Priority": "high",
            "Data": "full",
            "Compliance": "clear",
            "Compliance reasoning": "Unrelated retail analytics company based in Austria.",
            "Summary": "Acme builds retail dashboards on AWS.",
            "Size": "51-1000 (website)",
            "Cloud signal": "medium workload, inferred (website)",
            "HQ": "Austria",
            "Contact": "Ada Example, CTO <ada@acme.io>",
            "Updated": "2026-10-01 14:30",
        }
    ]


def test_fit_stays_a_number_so_the_sheet_can_be_sorted(path):
    record(make_result(), path)

    assert isinstance(rows(path)[0]["Fit"], int)


def test_another_company_is_appended(path):
    record(make_result(), path)
    row = record(make_result(company="Beta Corp", domain="beta.example"), path)

    assert row == 3
    assert [entry["Company"] for entry in rows(path)] == ["Acme Analytics", "Beta Corp"]


def test_the_same_company_updates_its_row_instead_of_adding_one(path):
    record(make_result(), path)
    record(make_result(company="Beta Corp", domain="beta.example"), path)

    row = record(make_result(route=Route.CHECK_FIRST, reason="Possible competitor"), path)

    assert row == 2
    assert len(rows(path)) == 2
    assert rows(path)[0]["Route"] == "CHECK-FIRST"


def test_another_company_on_the_same_domain_cannot_overwrite_a_row(path):
    """Found in review: anyone could replace a blocked company's row by submitting its website."""
    blocked = make_result(
        company="CloudTrim Inc", route=Route.DO_NOT_ENGAGE, reason="Competitor: CloudTrim Inc"
    )
    record(blocked, path)

    record(make_result(company="Victim GmbH"), path)

    assert {entry["Company"]: entry["Route"] for entry in rows(path)} == {
        "Victim GmbH": "SALES-READY",
        "CloudTrim Inc": "DO-NOT-ENGAGE",
    }


def test_the_sheet_is_kept_in_calling_order_with_blocked_leads_at_the_bottom(path):
    blocked = make_result(
        company="CloudTrim Inc", domain="ct.example", route=Route.DO_NOT_ENGAGE, fit=100
    )
    record(make_result(company="Medium Co", domain="medium.example", fit=45), path)
    record(blocked, path)
    record(
        make_result(company="Check Co", domain="c.example", route=Route.CHECK_FIRST, fit=70), path
    )
    row = record(make_result(company="Top Co", domain="top.example", fit=90), path)

    assert row == 2
    assert [(entry["Company"], entry["Fit"]) for entry in rows(path)] == [
        ("Top Co", 90),
        ("Check Co", 70),  # a lead with an open question keeps its place by score
        ("Medium Co", 45),
        ("CloudTrim Inc", 100),  # blocked: still visible, never above a callable lead
    ]


def test_the_same_company_is_recognised_despite_spelling_of_its_legal_form(path):
    record(make_result(company="Acme Analytics"), path)

    assert record(make_result(company="ACME Analytics, Inc."), path) == 2


def test_leads_without_a_website_are_matched_by_company_name(path):
    record(make_result(company="No Site Ltd", domain=None), path)
    record(make_result(company="Other No Site", domain=None), path)

    row = record(make_result(company="no site ltd", domain=None), path)

    assert row == 2
    assert len(rows(path)) == 2


@pytest.mark.parametrize("route", list(Route))
def test_the_route_cell_is_colour_coded(path, route):
    record(make_result(route=route), path)

    cell = load_workbook(path)["Leads"]["A2"]
    assert cell.fill.start_color.rgb.endswith(ROUTE_COLOURS[route])


@pytest.mark.parametrize(
    "payload", ['=HYPERLINK("http://evil.example","x")', "+1", "-1", "@SUM(A1)"]
)
def test_submitted_text_is_never_stored_as_a_formula(path, payload):
    record(make_result(company=payload, summary=payload), path)

    stored = rows(path)[0]
    assert stored["Company"] == f"'{payload}"
    assert stored["Summary"] == f"'{payload}"


def test_a_spreadsheet_open_in_excel_gives_a_clear_error(path, monkeypatch):
    def locked(_self, _path):
        raise PermissionError("locked")

    monkeypatch.setattr(Workbook, "save", locked)

    with pytest.raises(TrackerError, match="close it and run again"):
        record(make_result(), path)
